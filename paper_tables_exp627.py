"""Generate only the three EXP627 paper tables; never invoke experiment execution.

Run: .venv/bin/python -B paper_tables_exp627.py
Only Results/EXP627/full_rep1/Paper_Tables_EXP627.xlsx is replaced.
"""
from io import BytesIO
import inspect
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from full_replica_report import (
    CLASSIFIERS, DATASETS, TABLE_ORDER, framework, load_completed_full,
    report_guard, sha256,
    validate_replica_destination, recreate_replica_directories,
)


METRICS = (
    ("Accuracy", "AccRuns", "AS_test", "Accuracy", 100.0),
    ("F1-Score", "F1Runs", "F1_test", "F1Score", 1.0),
    ("Precision", "PSRuns", "PS_test", "Precision", 1.0),
    ("Recall", "RSRuns", "RS_test", "Recall", 1.0),
)
CLASSIFIER_ORDER = ("knn", "rf", "svm")
DATASET_LABELS = ("Data Class", "Feature Envy", "God Class", "Long Method",
                  "Long Parameter List", "Switch Statement")
STATS = ("Best", "Worst", "Mean", "Std")
SHEETS = ("Table2_Overall", "Table3_Datasets", "Table4_Datasets")


def protected_hashes(root, destination):
    files = set(root.rglob("*.pkl"))
    for kind in ("Results", "Figures"):
        files.update(p for p in (root / kind / "EXP627").rglob("*")
                     if p.is_file() and destination not in p.parents)
    return {str(p.relative_to(root)): sha256(p) for p in sorted(files)}


def require_close(actual, reference, label):
    if not np.allclose(actual, reference, rtol=1e-12, atol=1e-12):
        raise ValueError(f"EXP627 reference-mean validation failed: {label}")


def prepare_tables(args):
    """Validate report means first; calculate statistics only from original runs."""
    m = framework()
    args, results, indexed, _ = load_completed_full(args)
    classifier = inspect.signature(m._generate_seven_global_charts).parameters["estimator_filter"].default
    if classifier != "svm":
        raise ValueError("EXP627 per-dataset report classifier changed; audit required")
    source = Path(args.output_root) / "Results/EXP627/full"
    summary = pd.read_csv(source / "RESUMEN_GRAFICAS_EXP627.csv")
    key_columns = ["Dataset", "Estimator", "Optimizer"]
    if summary.duplicated(key_columns).any():
        raise ValueError("Ambiguous identities in original EXP627 report")
    summary = summary.set_index(key_columns)
    expected_keys = {(ds, cls, opt) for ds in DATASETS for cls in CLASSIFIERS for opt in TABLE_ORDER}
    if set(summary.index) != expected_keys:
        raise ValueError("Original EXP627 report must have all 216 unique combinations")
    existing = pd.read_excel(source / "Global_Results_EXP627.xlsx", sheet_name=None, index_col=0)
    # Reproduce every cached mean against two independently serialized original reports.
    for (dataset, cls, opt), row in indexed.items():
        for metric, run_key, csv_key, excel_sheet, scale in METRICS:
            mean = np.mean(row[run_key]) / scale
            require_close(mean, summary.loc[(dataset, cls, opt), csv_key], f"CSV/{dataset}/{cls}/{opt}/{metric}")
            require_close(mean, existing[excel_sheet].loc[dataset, f"{opt}_{cls.upper()}"] / scale,
                          f"Excel/{dataset}/{cls}/{opt}/{metric}")
    tables = {}
    overall = []
    overall_reference = []
    for opt in TABLE_ORDER:
        columns, mean_columns = [], []
        for metric, run_key, csv_key, _, scale in METRICS:
            for cls in CLASSIFIER_ORDER:
                # Position r is the same run number: execute_pending_runs writes in run order.
                runs = np.stack([indexed[ds, cls, opt][run_key] for ds in DATASETS]) / scale
                overall_runs = np.mean(runs, axis=0)
                assert overall_runs.shape == (30,)
                stats = m._run_stats(overall_runs, "max")
                reference = np.mean([summary.loc[(ds, cls, opt), csv_key] for ds in DATASETS])
                require_close(stats["Mean"], reference, f"Table2/{opt}/{cls}/{metric}")
                columns.append([stats[stat] for stat in STATS])
                mean_columns.append(reference)
        overall.extend(np.asarray(columns).T.tolist())
        overall_reference.append(mean_columns)
    tables[SHEETS[0]] = (np.asarray(overall), np.asarray(overall_reference))
    for sheet, datasets in zip(SHEETS[1:], (DATASETS[:3], DATASETS[3:])):
        rows, reference_rows = [], []
        for opt in TABLE_ORDER:
            columns, mean_columns = [], []
            for dataset in datasets:
                for metric, run_key, csv_key, _, scale in METRICS:
                    stats = m._run_stats(np.asarray(indexed[dataset, classifier, opt][run_key]) / scale, "max")
                    reference = summary.loc[(dataset, classifier, opt), csv_key]
                    require_close(stats["Mean"], reference, f"{sheet}/{dataset}/{opt}/{metric}")
                    columns.append([stats[stat] for stat in STATS])
                    mean_columns.append(reference)
            rows.extend(np.asarray(columns).T.tolist())
            reference_rows.append(mean_columns)
        tables[sheet] = (np.asarray(rows), np.asarray(reference_rows))
    return tables, classifier


def make_workbook(tables, classifier):
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = "EXP627 FULL — Paper Tables 2, 3 and 4"
    wb.properties.description = (
        "30 completed runs; sample Std (ddof=1); all metrics 0–1. "
        "Overall statistics use 30 per-run six-dataset averages. Dataset tables use SVM. "
        "Mean rows verified against original EXP627 Global Results and report CSV; "
        "manuscript screenshots are not available for direct comparison."
    )
    thin = Side(style="thin", color="DCE2E8")
    group_border = Side(style="medium", color="8394A5")
    for sheet_index, name in enumerate(SHEETS):
        ws = wb.create_sheet(name)
        values, references = tables[name]
        assert values.shape == (48, 12)
        ws.merge_cells("A1:B1")
        ws["A1"] = "Performance metric" if sheet_index == 0 else "Code Smell"
        ws["A2"], ws["B2"] = "Algorithm", "Statistic"
        if sheet_index == 0:
            for i, (metric, *_) in enumerate(METRICS):
                col = 3 + i*3
                ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col+2)
                ws.cell(1, col, metric)
                for j, cls in enumerate(CLASSIFIER_ORDER):
                    ws.cell(2, col+j, cls.upper())
            note = "Classifier subcolumns: KNN, RF, SVM. Each run is averaged across all six datasets before Best/Worst/Mean/Std."
        else:
            offset = (sheet_index-1)*3
            for i, dataset in enumerate(DATASET_LABELS[offset:offset+3]):
                col = 3 + i*4
                ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col+3)
                ws.cell(1, col, dataset)
                for j, (metric, *_) in enumerate(METRICS):
                    ws.cell(2, col+j, metric)
            note = f"Classifier: {classifier.upper()}, as selected by the existing EXP627 FULL per-dataset report."
        ws["A1"].comment = Comment(
            note + " All metrics use the existing report's 0–1 scale (cached Accuracy divided by 100). "
            "Exactly 30 runs; Std uses ddof=1. All Mean rows match the original Excel and CSV reports. "
            "Source: Results/EXP627/full/cache/*_1eccd45e74_results.pkl. Optimization calls: 0.",
            "EXP627 report validation",
        )
        for i, opt in enumerate(TABLE_ORDER):
            first = 3+i*4
            ws.merge_cells(start_row=first, start_column=1, end_row=first+3, end_column=1)
            ws.cell(first, 1, "DSA-DE" if opt == "DSADE" else opt)
            for stat_index, stat in enumerate(STATS):
                row = first+stat_index
                ws.cell(row, 2, stat)
                for col, value in enumerate(values[i*4+stat_index], 3):
                    ws.cell(row, col, float(value)).number_format = "0.0000"
                for cell in ws[row]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = Border(left=thin, right=thin, top=group_border if stat_index == 0 else thin,
                                         bottom=group_border if stat_index == 3 else thin)
                    cell.fill = PatternFill("solid", fgColor="F3F6FA" if i % 2 == 0 else "FFFFFF")
                    cell.font = Font(name="Calibri", size=11, bold=(opt == "DSADE"))
                ws.row_dimensions[row].height = 20
            ws.cell(first, 1).border = Border(left=thin, right=thin, top=group_border, bottom=group_border)
        for row in ws.iter_rows(min_row=1, max_row=2):
            for cell in row:
                cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
                cell.fill = PatternFill("solid", fgColor="294866")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = Border(bottom=thin, right=thin)
        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 25
        ws.column_dimensions["A"].width = 16
        ws.column_dimensions["B"].width = 12
        for col in range(3, 15):
            ws.column_dimensions[get_column_letter(col)].width = 12
        ws.freeze_panes = "C3"
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 85
        ws.print_title_rows = "1:2"
        ws.print_area = "A1:N50"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
        ws.print_options.horizontalCentered = True
        # Validate actual worksheet Mean rows before serializing.
        require_close([[ws.cell(5+i*4, col).value for col in range(3, 15)] for i in range(12)],
                      references, f"{name}/worksheet Mean rows")
    return wb


def run_paper_tables(args):
    root = Path(args.output_root).resolve()
    destination = root / "Results/EXP627/full_rep1"
    target = destination / "Paper_Tables_EXP627.xlsx"
    from full_rep1_report import validate_output_destination
    validate_output_destination(root, destination)
    before = protected_hashes(root, destination)
    # Reuse the report-only write/execution guard. No figure functions are called.
    with report_guard((destination,)) as guard:
        tables, classifier = prepare_tables(args)
        wb = make_workbook(tables, classifier)
        print("Validated: 6 datasets / 3 classifiers / 12 algorithms / 30 runs; SVM dataset tables.")
        print("All 432 Mean cells reproduce the original EXP627 report values (tolerance 1e-12).")
        destination.mkdir(exist_ok=True)
        import tempfile
        old_tempdir = tempfile.tempdir
        tempfile.tempdir = str(destination)
        try:
            buffer = BytesIO()
            wb.save(buffer)
        finally:
            tempfile.tempdir = old_tempdir
        buffer.seek(0)
        check = load_workbook(buffer, data_only=True)
        assert tuple(check.sheetnames) == SHEETS
        for name in SHEETS:
            ws = check[name]
            assert ws.max_row == 50 and ws.max_column == 14
            actual = [[ws.cell(row, col).value for col in range(3, 15)] for row in range(3, 51)]
            require_close(actual, tables[name][0], f"{name}/all serialized statistics")
            for i, opt in enumerate(TABLE_ORDER):
                first = 3+i*4
                assert f"A{first}:A{first+3}" in ws.merged_cells
                assert ws.cell(first, 1).value == ("DSA-DE" if opt == "DSADE" else opt)
                assert tuple(ws.cell(first+j, 2).value for j in range(4)) == STATS
            assert all(ws.cell(row, col).number_format == "0.0000" for row in range(3, 51) for col in range(3, 15))
        check.close()
        if protected_hashes(root, destination) != before:
            raise AssertionError("Protected original files changed before saving")
        target.write_bytes(buffer.getvalue())
        assert guard["optimization_calls"] == 0
    if protected_hashes(root, destination) != before:
        raise AssertionError("Protected original files changed")
    print(f"Created: {target}")
    print(f"Verified 3 worksheets / 1,728 numeric cells; optimization calls = 0; {len(before)} protected files unchanged (SHA-256).")


if __name__ == "__main__":
    run_paper_tables(framework().parse_args())
