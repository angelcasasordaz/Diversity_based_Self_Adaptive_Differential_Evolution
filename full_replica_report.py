"""Strict, additive EXP627 FULL reporting. No experiment execution path."""
import argparse
from contextlib import contextmanager
import hashlib
import json
import os
import shutil
from pathlib import Path
import sys

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break


DATASETS = ("DataClass", "FeatureEnvy", "GodClass", "LongMethod",
            "LongParameterList", "SwitchStatements")
CLASSIFIERS = ("knn", "svm", "rf")
CONFIG_ORDER = ("DSADE", "DE", "JADE", "SHADE", "PSO", "WOA",
                "HHO", "GOA", "SA", "BRO", "RUN", "FOX")
TABLE_ORDER = ("DSADE", "JADE", "BRO", "DE", "FOX", "GOA", "HHO",
               "PSO", "RUN", "SA", "SHADE", "WOA")
METRICS = (("Accuracy", "AccRuns"), ("F1-Score", "F1Runs"),
           ("Precision", "PSRuns"), ("Recall", "RSRuns"))
SIGNATURE = "1eccd45e74"
REPLICA_DIRECTORIES = (
    "Results/EXP627/full_replica",
    "Figures/EXP627/full_replica",
    "Results/EXP627/full_replica_tables",
    "Figures/EXP627/publication_unified",
)


def validate_replica_destination(root, destination):
    """Allow only exact replica roots; never follow links or remove scientific data."""
    root = Path(root).resolve()
    destination = Path(destination).absolute()
    allowed = {root / relative for relative in REPLICA_DIRECTORIES}
    if destination not in allowed or destination.resolve() != destination:
        raise ValueError(f"Unsafe replica deletion path: {destination}")
    if any(path.is_symlink() for path in (destination, *destination.parents)):
        raise ValueError(f"Symlink in replica deletion path: {destination}")
    if destination.is_symlink() or (destination.exists() and not destination.is_dir()):
        raise ValueError(f"Replica destination is not an ordinary directory: {destination}")
    if destination.exists():
        for parent, dirs, files in os.walk(destination, followlinks=False):
            for path in [Path(parent), *(Path(parent) / name for name in dirs + files)]:
                name = path.name.lower()
                if (path.is_symlink() or path.is_mount()
                        or path.suffix.lower() in {".pkl", ".pickle", ".ckpt"}
                        or name == "cache" or "checkpoint" in name or "progress" in name):
                    raise ValueError(f"Refusing replica deletion containing protected data/link: {path}")
    return destination


def recreate_replica_directories(root, *destinations):
    raise ValueError("Retired reporting directories cannot be recreated; use full_rep1")


def framework():
    # main_best.py may be running as __main__; reuse that exact module.
    main = sys.modules.get("__main__")
    if getattr(main, "__file__", "").endswith("main_best.py"):
        return main
    import main_best
    return main_best


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def original_hashes(root):
    paths = set(root.rglob("*.pkl"))
    for kind in ("Results", "Figures"):
        base = root / kind / "EXP627"
        paths.update(p for p in base.rglob("*")
                     if p.is_file() and not any(root / relative in p.parents for relative in REPLICA_DIRECTORIES))
    return {str(p.relative_to(root)): sha256(p) for p in sorted(paths)}


def load_completed_full(args):
    """Use existing identity/parser/validation logic, but only exact final caches."""
    m = framework()
    args = argparse.Namespace(**vars(args))
    if args.experiment_modes != ["full"] or args.exp_id != 627 or args.dataset_source != "codesmell":
        raise ValueError("REPORT-ONLY requires source EXP627 FULL / codesmell")
    args.experiment_mode = "full"
    m.apply_experiment_mode(args)
    if (tuple(m.optimizer_order_from_config(args.optimizers)) != CONFIG_ORDER
            or set(args.estimators) != set(CLASSIFIERS) or len(args.estimators) != 3
            or args.runs != 30 or args.transfer_functions != ["vstf_01"]):
        raise ValueError("REPORT-ONLY requires the original 12 optimizers, KNN/SVM/RF, vstf_01 and 30 runs")
    if args.datasets is not None and set(args.datasets) != set(DATASETS):
        raise ValueError("REPORT-ONLY requires all six EXP627 code-smell datasets")
    signature = m.build_cache_signature(args)
    if signature != SIGNATURE:
        raise ValueError(f"Incompatible EXP627 FULL configuration: {signature}; expected {SIGNATURE}. No recomputation allowed.")
    root = Path(args.output_root).resolve()
    results, indexed, sources = {}, {}, []
    for dataset in DATASETS:
        results[dataset] = {}
        for classifier in CLASSIFIERS:
            path = root / "Results/EXP627/full/cache" / f"EXP627_{dataset}_{classifier}_{signature}_results.pkl"
            if not path.is_file():
                raise FileNotFoundError(f"Missing completed EXP627 FULL cache: {path}. No recomputation allowed.")
            payload = m.load_cache(str(path))
            expected = m.expected_result_labels(args, classifier, False, True)
            if not isinstance(payload, dict) or len(payload) != 12:
                raise ValueError(f"Expected exactly 12 optimizer rows in {path}")
            for label, legacy in expected:
                matches = [key for key in dict.fromkeys((label, legacy)) if key in payload]
                if len(matches) != 1:
                    raise ValueError(f"Missing/ambiguous optimizer {label} in {path}")
                source_label = matches[0]
                row = payload[source_label]
                values, reason = m.validate_source_label_runs(row, classifier, 30)
                if reason or row.get("CompletedRuns") != 30:
                    raise ValueError(f"Incomplete {dataset}/{label}: {reason or 'requires 30 completed runs'}")
                for field, items in values.items():
                    arr = np.asarray(items, dtype=float)
                    shape = (30, args.epochs) if field == "CurvesAll" else (30,)
                    if arr.shape != shape or not np.isfinite(arr).all():
                        raise ValueError(f"Invalid/missing run values: {dataset}/{label}/{field}, expected {shape}")
                for run_key, mean_key in (("AccRuns", "AccMean"), ("F1Runs", "F1Mean"),
                                          ("PSRuns", "PSMean"), ("RSRuns", "RSMean"),
                                          ("FitRuns", "FitMean"), ("FeatRuns", "FeatMean"),
                                          ("TimeRuns", "TimeMean")):
                    if not np.isclose(row.get(mean_key, np.nan), np.mean(values[run_key]), rtol=1e-12, atol=1e-12):
                        raise ValueError(f"Cached summary disagrees with runs: {dataset}/{label}/{mean_key}")
                curve = np.asarray(row.get("Curve", []), dtype=float)
                if curve.shape != (args.epochs,) or not np.allclose(curve, m.pad_mean_curves(values["CurvesAll"], args.epochs), rtol=1e-12, atol=1e-12):
                    raise ValueError(f"Invalid cached mean convergence: {dataset}/{label}")
                parsed = m.parse_result_label(source_label, args)
                key = (dataset, parsed["estimator"], parsed["method"])
                if key != (dataset, classifier, m.parse_result_label(label, args)["method"]) or key in indexed:
                    raise ValueError(f"Ambiguous scientific identity: {source_label}")
                indexed[key] = row
                results[dataset][label] = row
            sources.append(str(path.relative_to(root)))
    return args, results, indexed, sources


def export_extended_statistical_excel(indexed, out_path):
    """Classifier-specific manuscript tables; retain original cached units."""
    m = framework()
    wb = Workbook()
    wb.remove(wb.active)
    checks = []
    thin = Side(style="thin", color="D9E1E8")
    for classifier in CLASSIFIERS:
        ws = wb.create_sheet(classifier.upper())
        ws.freeze_panes = "B5"
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 16
        for col in range(2, 14):
            ws.column_dimensions[get_column_letter(col)].width = 13
        ws.merge_cells("A1:M1")
        ws["A1"] = f"EXP627 FULL — {classifier.upper()} — 30 runs; Accuracy (%), other metrics (0–1); sample Std (ddof=1)"
        ws["A1"].font = Font(bold=True, size=12)
        rownum = 3
        for statistic in ("Best", "Worst", "Mean", "Std"):
            if rownum > 3:
                ws.row_breaks.append(Break(id=rownum - 1))
            ws.merge_cells(start_row=rownum, start_column=1, end_row=rownum, end_column=13)
            title = ws.cell(rownum, 1, statistic.upper())
            title.font = Font(bold=True, color="FFFFFF", size=12)
            title.fill = PatternFill("solid", fgColor="17365D")
            ws.row_dimensions[rownum].height = 23
            rownum += 1
            for datasets in (DATASETS[:3], DATASETS[3:]):
                ws.merge_cells(start_row=rownum, start_column=1, end_row=rownum+1, end_column=1)
                ws.cell(rownum, 1, "Algorithm")
                for i, dataset in enumerate(datasets):
                    first = 2 + 4*i
                    ws.merge_cells(start_row=rownum, start_column=first, end_row=rownum, end_column=first+3)
                    ws.cell(rownum, first, dataset)
                    for j, (metric, _) in enumerate(METRICS):
                        ws.cell(rownum+1, first+j, metric)
                for cells in ws.iter_rows(min_row=rownum, max_row=rownum+1, max_col=13):
                    for cell in cells:
                        cell.font = Font(bold=True, color="17365D")
                        cell.fill = PatternFill("solid", fgColor="DCE6F1")
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.border = Border(bottom=thin)
                for i, optimizer in enumerate(TABLE_ORDER):
                    r = rownum+2+i
                    ws.cell(r, 1, "DSA-DE" if optimizer == "DSADE" else optimizer)
                    for j, dataset in enumerate(datasets):
                        cached = indexed[dataset, classifier, optimizer]
                        for k, (_, run_key) in enumerate(METRICS):
                            value = m._run_stats(cached[run_key], "max")[statistic]
                            cell = ws.cell(r, 2+4*j+k, value)
                            cell.number_format = "0.0000"
                            checks.append((ws.title, cell.coordinate, value))
                    for cell in ws[r]:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.border = Border(bottom=thin)
                        if optimizer == "DSADE":
                            cell.font = Font(bold=True, color="0072B2")
                        if i % 2 == 0:
                            cell.fill = PatternFill("solid", fgColor="F2F6FA")
                rownum += 16
            rownum += 2
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_options.horizontalCentered = True
        ws.print_area = f"A1:M{rownum-3}"
    with Path(out_path).open("xb") as stream:
        wb.save(stream)
    # Read the actual XLSX back and verify every one of its 3,456 numeric cells.
    check = load_workbook(out_path, data_only=True)
    for sheet, coordinate, expected in checks:
        cell = check[sheet][coordinate]
        if not np.isclose(cell.value, expected, rtol=1e-14, atol=1e-14) or cell.number_format != "0.0000":
            raise AssertionError(f"Workbook verification failed: {sheet}!{coordinate}")
    check.close()
    return len(checks)


@contextmanager
def report_guard(destinations):
    """Reject scientific execution and filesystem mutations outside replica folders."""
    state = {"active": True, "optimization_calls": 0}
    forbidden = {"_run_single", "execute_pending_runs", "build_optimizer",
                 "configure_compute_backend", "save_cache", "start_gpu_request_service"}
    def profile(frame, event, arg):
        if event == "call" and (frame.f_code.co_name in forbidden or
                                (frame.f_code.co_name == "solve" and "mealpy" in frame.f_code.co_filename) or
                                (frame.f_code.co_name == "fit" and "mafese" in frame.f_code.co_filename)):
            state["optimization_calls"] += 1
            raise RuntimeError(f"REPORT-ONLY blocked scientific execution: {frame.f_code.co_name}")
    def allowed(path):
        if isinstance(path, int):
            return False
        resolved = Path(os.fsdecode(path)).resolve()
        return any(resolved == base or base in resolved.parents for base in destinations) and resolved.suffix != ".pkl"
    def audit(event, args):
        if not state["active"]:
            return
        targets = []
        def at(path, dir_fd):
            # shutil.rmtree uses descriptor-relative unlink/rmdir on Linux.
            if dir_fd is not None and dir_fd != -1 and not os.path.isabs(path):
                return Path(os.readlink(f"/proc/self/fd/{dir_fd}")) / os.fsdecode(path)
            return path
        if event == "open":
            path, mode, flags = args
            if (flags & (os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_TRUNC | os.O_APPEND)):
                targets = [path]
        elif event in {"os.remove", "os.rmdir"}:
            targets = [at(args[0], args[1])]
        elif event in {"os.mkdir", "os.chmod"}:
            targets = [at(args[0], args[2])]
        elif event in {"os.utime", "os.truncate", "shutil.rmtree"}:
            targets = [args[0]]
        elif event in {"os.rename", "os.link", "os.symlink"}:
            targets = list(args[:2])
        elif event in {"subprocess.Popen", "os.system", "os.fork"}:
            raise RuntimeError(f"REPORT-ONLY blocked process execution: {event}")
        if any(not allowed(path) for path in targets):
            raise PermissionError(f"REPORT-ONLY blocked mutation outside replica output: {event} {targets}")
    previous = sys.getprofile()
    sys.addaudithook(audit)
    sys.setprofile(profile)
    try:
        yield state
    finally:
        state["active"] = False
        sys.setprofile(previous)


def run_full_replica_report(args):
    """Compatibility entry point for the consolidated eight-figure report."""
    from full_rep1_report import run
    return run(args)
