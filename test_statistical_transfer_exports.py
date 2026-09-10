"""Read-only regressions using saved EXP626 runs; never execute optimization."""
import argparse
import hashlib
import pickle
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
from openpyxl import load_workbook

import main_best as study


class SavedTransferStatisticsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory = Path(__file__).parent / 'Results/EXP626/transfer_functions'
        cls.datasets = list(study.TRANSFER_FUNCTION_DATASETS)
        cls.signature = '0fbda34b93'
        cls.paths = study.Paths('EXP626', 'transfer_functions', '', str(cls.directory),
                                str(cls.directory / 'cache'))
        cls.args = argparse.Namespace(
            experiment_mode='transfer_functions', optimizers=['MaCRO-DE'],
            estimators=['knn', 'svm'], transfer_functions=list(study.TRANSFER_FUNCTION_TESTS),
        )
        cls.saved = {}
        for dataset in cls.datasets:
            for estimator in cls.args.estimators:
                path = cls.directory / 'cache' / f'EXP626_{dataset}_{estimator}_{cls.signature}_results.pkl'
                if not path.exists():
                    raise unittest.SkipTest('Saved EXP626 transfer-function caches are required')
                cls.saved[dataset, estimator] = pickle.loads(path.read_bytes())
        cls.hashes = {p: hashlib.sha256(p.read_bytes()).hexdigest()
                      for p in (cls.directory / 'cache').glob('*.pkl')}

    @classmethod
    def tearDownClass(cls):
        assert all(hashlib.sha256(p.read_bytes()).hexdigest() == digest
                   for p, digest in cls.hashes.items()), 'Cache changed'

    def setUp(self):
        guard = patch.object(study, 'execute_pending_runs', side_effect=AssertionError('Optimization forbidden'))
        guard.start()
        self.addCleanup(guard.stop)

    def test_all_variants_classifiers_metrics_and_order(self):
        results = study.load_results_from_cache(
            self.paths, self.args, self.datasets, self.signature, preserve_classifier_labels=True,
        )
        for dataset in self.datasets:
            self.assertEqual(len(results[dataset]), sum(len(self.saved[dataset, est]) for est in self.args.estimators))
        metrics = {
            'Accuracy': ('AccRuns', True), 'Precision': ('PSRuns', True),
            'Recall': ('RSRuns', True), 'F1Score': ('F1Runs', True),
            'Fitness': ('FitRuns', False), 'Features': ('FeatRuns', False),
            'Time': ('TimeRuns', False),
        }
        with tempfile.TemporaryDirectory() as tmp:
            for estimators in [['knn', 'svm'], ['knn'], ['svm']]:
                # Reversing the order verifies that the exporter honors configuration.
                transfers = list(reversed(self.args.transfer_functions))
                args = argparse.Namespace(**vars(self.args))
                args.estimators = estimators
                args.transfer_functions = transfers
                subset = {ds: {label: row for label, row in data.items()
                               if row['Estimator'].lower() in estimators}
                          for ds, data in results.items()}
                path = Path(tmp) / 'statistics.xlsx'
                study.export_statistical_excel(subset, self.datasets, args.optimizers, args, str(path))
                workbook = load_workbook(path)
                self.assertEqual(workbook.sheetnames, list(metrics))
                headers = [f'{est.upper()} | {tf.upper()}' if len(estimators) > 1 else tf.upper()
                           for est in estimators for tf in transfers]
                for sheet, (key, maximize) in metrics.items():
                    ws = workbook[sheet]
                    self.assertEqual([c.value for c in ws[1]], ['Dataset', 'Statistic'] + headers)
                    self.assertEqual({str(r) for r in ws.merged_cells.ranges},
                                     {f'A{2+4*i}:A{5+4*i}' for i in range(len(self.datasets))})
                    for i, dataset in enumerate(self.datasets):
                        self.assertEqual(ws.cell(2+4*i, 1).value, dataset)
                        self.assertEqual([ws.cell(2+4*i+j, 2).value for j in range(4)],
                                         ['Best', 'Worst', 'Mean', 'Std'])
                        for column, (est, tf) in enumerate(
                            ((est, tf) for est in estimators for tf in transfers), start=3
                        ):
                            values = np.asarray(self.saved[dataset, est][f'MACRO-DE_{tf.upper()}'][key], dtype=float)
                            values = values[np.isfinite(values)]
                            expected = [values.max() if maximize else values.min(),
                                        values.min() if maximize else values.max(),
                                        values.mean(), values.std(ddof=1) if len(values) > 1 else 0.0]
                            for j, value in enumerate(expected):
                                self.assertEqual(ws.cell(2+4*i+j, column).value, float(format(value, '.16g')))

    def test_available_variants_outside_configuration_are_retained(self):
        results = study.load_results_from_cache(
            self.paths, self.args, self.datasets, self.signature, preserve_classifier_labels=True,
        )
        args = argparse.Namespace(**vars(self.args))
        args.transfer_functions = ['sstf_04']
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'statistics.xlsx'
            study.export_statistical_excel(results, self.datasets, args.optimizers, args, str(path))
            headers = [cell.value for cell in load_workbook(path)['Accuracy'][1]][2:]
            self.assertEqual(headers[0], 'KNN | SSTF_04')
            self.assertEqual(len(headers), 16)
            self.assertEqual(set(headers), {f'{est.upper()} | {tf.upper()}'
                                           for est in self.args.estimators for tf in self.args.transfer_functions})

    def test_cached_output_routes_complete_records_only_to_statistics(self):
        with patch.object(study, 'export_mode_outputs') as export:
            study.regenerate_figures_from_cache(self.paths, self.args, self.datasets, self.signature)
        call = export.call_args
        for dataset in self.datasets:
            self.assertEqual(len(call.kwargs['statistical_results'][dataset]), 16)
            self.assertEqual(len(call.args[3][dataset]), 8)


if __name__ == '__main__':
    unittest.main()
