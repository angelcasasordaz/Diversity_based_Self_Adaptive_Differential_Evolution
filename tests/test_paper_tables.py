"""Small synthetic manuscript workbooks; no experiments or figure rendering."""
from contextlib import ExitStack
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import main_best as m
from reporting import paper_tables as paper
from reporting import exp627_excel_tables as preset


def fixture(datasets=("First", "Second"), algorithms=("DE", "PSO"), classifiers=("knn", "rf"), runs=3):
    args = SimpleNamespace(experiment_mode="full", estimators=list(classifiers),
                           optimizers=list(algorithms), transfer_functions=["vstf_01"],
                           full_replica_report_only=False, exp_id=913)
    results = {}
    for d, ds in enumerate(datasets):
        results[ds] = {}
        for a, opt in enumerate(algorithms):
            for c, cls in enumerate(classifiers):
                n = runs if isinstance(runs, int) else runs[a, c]
                base = np.arange(n, dtype=float) + d*2 + a + c
                results[ds][f"{opt}_{cls.upper()}"] = {
                    "Estimator": cls, "AccRuns": 60+base*2, "F1Runs": .5+base/100,
                    "PSRuns": .6+base/100, "RSRuns": .7+base/100,
                    "FitRuns": .3+base/100, "FeatRuns": 2+base, "TimeRuns": 1+base,
                }
    return args, results


class PaperTableTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.paths = SimpleNamespace(exp_tag="EXP913", res_dir=self.temp.name, fig_dir=self.temp.name)

    def export(self, args, results):
        path = paper.export_paper_tables(results, list(results), args.optimizers, args, self.paths)
        self.assertEqual(Path(path).name, "Paper_Tables_EXP913.xlsx")
        wb = load_workbook(path, data_only=True)
        self.addCleanup(wb.close)
        return wb

    def test_all_statistics_units_directions_and_per_run_averaging(self):
        args, results = fixture()
        wb = self.export(args, results)
        self.assertEqual(wb.sheetnames, ["Overall", "Datasets_1", "Datasets_2"])
        for a, opt in enumerate(args.optimizers):
            for metric_index, metric in enumerate(paper.METRICS):
                for c, cls in enumerate(args.estimators):
                    vectors = [np.asarray(records[f"{opt}_{cls.upper()}"][metric.run_key])/metric.scale
                               for records in results.values()]
                    overall = np.mean(vectors, axis=0)
                    expected = [max(overall), min(overall), np.mean(overall), np.std(overall, ddof=1)]
                    if metric.best_mode == "min": expected[:2] = expected[1::-1]
                    actual = [wb["Overall"].cell(3+4*a+s, 3+2*metric_index+c).value for s in range(4)]
                    np.testing.assert_allclose(actual, expected, rtol=1e-14)
                    for d, vector in enumerate(vectors):
                        expected = [max(vector), min(vector), np.mean(vector), np.std(vector, ddof=1)]
                        if metric.best_mode == "min": expected[:2] = expected[1::-1]
                        actual = [wb[f"Datasets_{c+1}"].cell(3+4*a+s, 3+7*d+metric_index).value for s in range(4)]
                        np.testing.assert_allclose(actual, expected, rtol=1e-14)
        self.assertEqual(wb["Overall"].freeze_panes, "C3")
        self.assertEqual(wb["Overall"]["C3"].number_format, "0.0000")

    def test_variable_dimensions_classifiers_metrics_and_run_counts(self):
        for datasets, algorithms, classifiers, runs in (
            (("Only",), ("DE",), ("tree",), 1),
            (("A", "B", "C", "D"), ("DE", "PSO", "JADE"), ("knn", "rf"),
             {(a, c): a+c+1 for a in range(3) for c in range(2)}),
        ):
            with self.subTest(datasets=datasets):
                args, results = fixture(datasets, algorithms, classifiers, runs)
                for records in results.values():
                    for row in records.values():
                        for metric in paper.METRICS[1:]: row.pop(metric.run_key)
                wb = self.export(args, results)
                self.assertEqual(wb["Overall"].max_row, len(algorithms)*4+2)
                self.assertEqual(wb["Overall"].max_column, len(classifiers)+2)
                self.assertEqual(len(wb.sheetnames), len(classifiers)+1)
                self.assertEqual(wb["Overall"]["C6"].value, 0)
                self.assertIn("run counts", wb.properties.description)

    def test_nonfinite_values_use_existing_reducer(self):
        args, results = fixture(datasets=("Only",), algorithms=("DE",), classifiers=("knn",))
        results["Only"]["DE_KNN"]["AccRuns"] = [np.nan, 80, np.inf]
        results["Only"]["DE_KNN"]["F1Runs"] = [np.nan]*3
        wb = self.export(args, results)
        self.assertEqual([wb["Overall"].cell(r, 3).value for r in range(3, 7)], [.8, .8, .8, 0])
        self.assertEqual([wb["Overall"].cell(r, 4).value for r in range(3, 7)], [None]*4)

    def test_incomplete_or_ambiguous_runs_are_rejected_without_writing(self):
        for problem in ("unequal", "missing", "duplicate", "empty"):
            args, results = fixture()
            row = results["First"]["DE_KNN"]
            if problem == "unequal": row["AccRuns"] = [80]
            if problem == "missing": del row["AccRuns"]
            if problem == "duplicate": results["First"]["DE_VSTF_01_KNN"] = row
            if problem == "empty": row["AccRuns"] = []
            with self.subTest(problem=problem), self.assertRaises(ValueError):
                self.export(args, results)
        self.assertFalse(list(Path(self.temp.name).glob("*.xlsx")))

    def test_full_pipeline_adds_workbook_alongside_existing_exports(self):
        args, results = fixture()
        with patch.object(m, "generate_seven_global_charts", return_value=[]) as charts:
            exported, _, _, statistical, friedman = m.export_mode_outputs(self.paths, args, list(results), results)
        self.assertEqual({Path(p).name for p in exported}, {"Global_Results_EXP913.xlsx", "Paper_Tables_EXP913.xlsx"})
        self.assertTrue(Path(statistical).is_file())
        self.assertTrue(Path(friedman).is_file())
        charts.assert_called_once()
        # Compare all values of the original workbooks with independent direct exports.
        direct = Path(self.temp.name)/"direct.xlsx"
        for path, exporter, params in (
            (exported[0], m.export_global_excel, (results, list(results))),
            (statistical, m.export_statistical_excel, (results, list(results), args.optimizers, args)),
            (friedman, m.export_friedman_analysis, (results, list(results), args.optimizers, args)),
        ):
            exporter(*params, str(direct))
            a, b = load_workbook(path, data_only=True), load_workbook(direct, data_only=True)
            try:
                self.assertEqual(a.sheetnames, b.sheetnames)
                for name in a.sheetnames:
                    self.assertEqual(list(a[name].values), list(b[name].values))
            finally:
                a.close(); b.close()

    def test_other_modes_and_report_only_do_not_invoke_exporter(self):
        class StopAfterExcel(Exception): pass
        for mode in ("full", "ablation", "sensitivity", "sensitivity_weights", "transfer_functions"):
            args, results = fixture()
            args.experiment_mode = mode
            args.full_replica_report_only = mode == "full"
            with ExitStack() as stack:
                for name in ("export_global_excel", "export_statistical_excel", "export_friedman_analysis",
                             "experiment_output_prefix"):
                    stack.enter_context(patch.object(m, name))
                stack.enter_context(patch.object(m, "generate_summary_dataframe", side_effect=StopAfterExcel))
                exporter = stack.enter_context(patch.object(paper, "export_paper_tables"))
                with self.assertRaises(StopAfterExcel):
                    m.export_mode_outputs(self.paths, args, list(results), results)
                exporter.assert_not_called()
            with self.assertRaises(ValueError):
                self.export(args, results)

    def test_exp627_preset_layout_and_all_scientific_values(self):
        args, results = fixture(preset.DATASETS, preset.TABLE_ORDER, preset.CLASSIFIERS, 30)
        args.output_root = self.temp.name
        indexed = {(ds, row["Estimator"], label.rsplit("_", 1)[0]): row
                   for ds, records in results.items() for label, row in records.items()}
        summary = pd.DataFrame([
            dict(Dataset=ds, Estimator=cls, Optimizer=opt,
                 **{csv_key: np.mean(row[run_key])/scale for _, run_key, csv_key, _, scale in preset.METRICS})
            for (ds, cls, opt), row in indexed.items()])
        existing = {sheet: pd.DataFrame({f"{opt}_{cls.upper()}":
                    [np.mean(indexed[ds, cls, opt][key]) for ds in preset.DATASETS]
                    for opt in preset.TABLE_ORDER for cls in preset.CLASSIFIERS}, index=preset.DATASETS)
                    for _, key, _, sheet, _ in preset.METRICS}
        with patch.object(preset, "load_completed_full", return_value=(args, results, indexed, [])), \
                patch.object(preset.pd, "read_csv", return_value=summary), \
                patch.object(preset.pd, "read_excel", return_value=existing):
            tables, classifier = preset.prepare_tables(args)
            (Path(self.temp.name)/"Results/EXP627").mkdir(parents=True)
            preset.run_paper_tables(args)
            self.assertTrue((Path(self.temp.name)/"Results/EXP627/full_rep1/Paper_Tables_EXP627.xlsx").is_file())
        wb = load_workbook(BytesIO(paper.workbook_bytes(preset.make_workbook(tables, classifier), tables)), data_only=True)
        self.addCleanup(wb.close)
        self.assertEqual(wb.sheetnames, list(preset.SHEETS))
        for sheet_index, name in enumerate(preset.SHEETS):
            ws = wb[name]
            self.assertEqual((ws.max_row, ws.max_column), (50, 14))
            self.assertEqual(ws["A3"].value, "DSA-DE")
            self.assertIn("A3:A6", ws.merged_cells)
            for a, opt in enumerate(preset.TABLE_ORDER):
                vectors = []
                if sheet_index == 0:
                    for _, key, _, _, scale in preset.METRICS:
                        for cls in preset.CLASSIFIER_ORDER:
                            vectors.append(np.mean([indexed[ds, cls, opt][key] for ds in preset.DATASETS], axis=0)/scale)
                else:
                    for ds in preset.DATASETS[(sheet_index-1)*3:sheet_index*3]:
                        for _, key, _, _, scale in preset.METRICS:
                            vectors.append(np.asarray(indexed[ds, "svm", opt][key])/scale)
                for c, vector in enumerate(vectors, 3):
                    expected = [max(vector), min(vector), np.mean(vector), np.std(vector, ddof=1)]
                    np.testing.assert_allclose([ws.cell(3+a*4+s, c).value for s in range(4)], expected, rtol=1e-14)


if __name__ == "__main__":
    unittest.main()
