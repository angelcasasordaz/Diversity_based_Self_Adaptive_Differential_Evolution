"""Read-only regression against saved scientific values; never render historical figures."""
from contextlib import ExitStack
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.stats import rankdata, t

import main_best as m
from reporting import core, figures, statistics, paper_tables


def exact_signed_rank_check(differences):
    """Independent subset-sum test oracle for nonzero, untied signed ranks."""
    d = np.asarray(differences)
    assert np.all(d != 0) and len(np.unique(abs(d))) == len(d)
    ranks = rankdata(abs(d)).astype(int)
    positive, total = int(ranks[d > 0].sum()), int(ranks.sum())
    counts = np.zeros(total + 1, dtype=np.int64)
    counts[0] = 1
    for rank in ranks:
        counts[rank:] += counts[:-rank].copy()
    w = min(positive, total - positive)
    return w, min(1., 2 * counts[:w+1].sum() / (2**len(d)))


class HistoricalScientificValuesTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = Path(__file__).resolve().parents[1]
        cls.source = cls.root / 'Results/EXP627/full'
        if not (cls.source / 'Paper_Tables_EXP627.xlsx').is_file():
            raise unittest.SkipTest('Saved EXP627 regression workbooks unavailable')
        with patch.object(sys, 'argv', ['main_best.py', '--full-replica-report-only', '--exp-id', '627',
                                       '--experiment-mode', 'full', '--output-root', str(cls.root)]):
            args = m.parse_args()
        # EXP-specific configuration lives only in this regression fixture.
        args.datasets = ['DataClass', 'FeatureEnvy', 'GodClass', 'LongMethod', 'LongParameterList', 'SwitchStatements']
        args.optimizers = ['DSADE', 'DE', 'JADE', 'SHADE', 'PSO', 'WOA', 'HHO', 'GOA', 'SA', 'BRO', 'RUN', 'FOX']
        args.estimators = ['knn', 'svm', 'rf']
        args.runs = 30
        args.transfer_functions = ['vstf_01']
        with core.report_guard(()):
            cls.report = core.load_completed_cache(args)
        cls.hashes = {p: core.sha256(p) for kind in ('Results', 'Figures')
                      for p in (cls.root / kind / 'EXP627').rglob('*') if p.is_file()}

    @classmethod
    def tearDownClass(cls):
        if any(core.sha256(path) != digest for path, digest in cls.hashes.items()):
            raise AssertionError('Historical cache/report changed')

    def setUp(self):
        stack = ExitStack(); self.addCleanup(stack.close)
        for name in ('_run_single', 'execute_pending_runs', 'build_optimizer', 'save_cache', 'load_dataset'):
            blocked = stack.enter_context(patch.object(m, name, side_effect=AssertionError('Scientific execution forbidden')))
            self.addCleanup(blocked.assert_not_called)

    def test_generic_means_ci_and_radar_preserve_historical_values(self):
        reference = pd.read_csv(self.source / 'RESUMEN_GRAFICAS_EXP627.csv')
        report = self.report
        columns = {'AccRuns': 'AS_test', 'PSRuns': 'PS_test', 'RSRuns': 'RS_test',
                   'F1Runs': 'F1_test', 'FeatRuns': 'N_Features_Selected', 'TimeRuns': 'Runtime'}
        for classifier in report.classifiers:
            for metric in report.metrics:
                if metric.run_key not in columns:
                    continue
                expected = figures.metric_values(reference, columns[metric.run_key], classifier,
                                                   report.datasets, report.algorithms)
                np.testing.assert_allclose(figures.metric_matrix(report, classifier, metric), expected, rtol=1e-12, atol=1e-12)
            precision = next(metric for metric in report.metrics if metric.run_key == 'PSRuns')
            values = figures.metric_matrix(report, classifier, precision)
            means, ci = figures.dataset_mean_ci(values)
            np.testing.assert_allclose(means, values.mean(axis=1), rtol=1e-14)
            np.testing.assert_allclose(ci, t.ppf(.975, 5) * values.std(axis=1, ddof=1) / np.sqrt(6), rtol=1e-14)
            labels, radar = figures.radar_values(report, classifier)
            self.assertEqual(len(labels), 5)
            features = figures.metric_values(reference, 'N_Features_Selected', classifier, report.datasets, report.algorithms)
            np.testing.assert_allclose(radar[:, :, -1], 1 - features / np.maximum(features.max(axis=0), 1), rtol=1e-12)

    def test_generic_statistics_preserve_saved_summary_and_all_pairs(self):
        report = self.report
        metric = next(metric for metric in report.metrics if metric.run_key == 'F1Runs')
        blocks, x = statistics.matched_block_matrix(report.indexed, report.datasets, report.classifiers, report.algorithms, metric)
        analysis = statistics.analyze(x, report.algorithms)
        source = self.root / 'Results/EXP627/full_rep1/statistics'
        reference = pd.read_csv(source / 'statistical_summary.csv').replace({'Algorithm': {'DSA-DE': 'DSADE'}})
        actual = analysis['summary'].rename(columns={'Mean': 'Mean_F1', 'Std': 'Std_F1', 'Median': 'Median_F1'})
        pd.testing.assert_frame_equal(actual, reference, check_exact=False, rtol=1e-12, atol=1e-14)
        reference = pd.read_csv(source / 'pairwise_wilcoxon_holm.csv').replace(
            {'Algorithm_A': {'DSA-DE': 'DSADE'}, 'Algorithm_B': {'DSA-DE': 'DSADE'}})
        actual = analysis['pairs'][reference.columns]
        pd.testing.assert_frame_equal(actual, reference, check_dtype=False, check_exact=False, rtol=1e-12, atol=1e-14)
        self.assertEqual(len(blocks), 18)
        self.assertEqual(len(actual), 66)
        for row in actual.itertuples():
            i, j = report.algorithms.index(row.Algorithm_A), report.algorithms.index(row.Algorithm_B)
            w, p = exact_signed_rank_check(x[:, i] - x[:, j])
            self.assertEqual(row.Wilcoxon_statistic, w)
            self.assertAlmostEqual(row.Raw_p, p, places=14)

    def test_generic_paper_tables_preserve_every_saved_numeric_cell(self):
        report = self.report
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'Paper_Tables_reference.xlsx'
            paper_tables.export_indexed_tables(report.indexed, report.datasets, report.algorithms,
                                               report.classifiers, report.metrics, path, title='Regression reference')
            actual = load_workbook(path, data_only=True)
            reference = load_workbook(self.source / 'Paper_Tables_EXP627.xlsx', data_only=True)
            try:
                self.assertEqual(actual.sheetnames, reference.sheetnames)
                paper_tables.validate_plain_workbook(actual)
                for name in reference.sheetnames:
                    a, b = list(actual[name].values), list(reference[name].values)
                    self.assertEqual((actual[name].max_row, actual[name].max_column),
                                     (reference[name].max_row, reference[name].max_column))
                    for ar, br in zip(a, b):
                        for av, bv in zip(ar, br):
                            if isinstance(bv, (int, float)):
                                np.testing.assert_allclose(av, bv, rtol=1e-12, atol=1e-12)
                            else:
                                self.assertEqual(av, bv)
            finally:
                actual.close(); reference.close()

    def test_exp627_report_only_completes_in_temporary_root(self):
        from tests.test_generic_reporting import tiny_figures, tiny_base_figures
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            shutil.copytree(self.source / 'cache', root / 'Results/EXP627/full/cache')
            args = self.report.args
            argv = ['main_best.py', '--full-replica-report-only', '--exp-id', '627',
                    '--experiment-mode', 'full', '--output-root', folder,
                    '--datasets', *self.report.datasets, '--optimizers', *args.optimizers,
                    '--estimators', *args.estimators, '--runs', str(args.runs),
                    '--epochs', str(args.epochs), '--transfer-functions', *args.transfer_functions]
            # Exercise the actual CLI, Excel exporters, statistics and validation;
            # only figure construction is replaced to avoid publication rendering.
            with patch.object(sys, 'argv', argv), \
                    patch.object(figures, 'base_publication_figures', tiny_base_figures), \
                    patch.object(figures, 'publication_figures', tiny_figures), \
                    patch.object(figures, 'statistical_figures', tiny_figures):
                manifest = m.main()
            self.assertEqual(manifest['experiment_id'], 627)
            self.assertEqual(manifest['optimization_calls'], 0)
            self.assertTrue(manifest['protected_files_unchanged'])
            self.assertTrue((root / 'Results/EXP627/full_rep1/Paper_Tables_EXP627.xlsx').is_file())
            self.assertFalse(list(root.rglob('*.pdf')))
