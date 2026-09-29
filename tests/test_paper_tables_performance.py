"""Merged-cell validation regressions, without timing-sensitive assertions."""
from io import BytesIO
import hashlib
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.merge import MergedCellRange

from reporting import paper_tables as paper


def representative_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Many merged ranges"
    for row in ws.iter_rows(min_row=1, max_row=162, min_col=1, max_col=44):
        for cell in row:
            cell.value = cell.row / 100 + cell.column / 1000
            cell.number_format = "0.0000"
    for row in range(3, 163, 4):
        ws.merge_cells(start_row=row, start_column=1, end_row=row+3, end_column=1)
        ws.cell(row, 1, f"Algorithm {row}")
    for col in range(3, 45, 3):
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col+2)
        ws.cell(1, col, f"Metric {col}")
    ws.merge_cells("A1:B1")
    ws["A1"] = "Performance"
    ws["C3"] = "=SUM(D3:E3)"
    return wb


def workbook_parts(wb):
    stream = BytesIO()
    wb.save(stream)
    with ZipFile(stream) as archive:
        # Saving updates the modification timestamp independently of validation.
        return {name: archive.read(name) for name in archive.namelist() if name != "docProps/core.xml"}


class MergedCellValidationTests(unittest.TestCase):
    def test_large_workbook_uses_one_index_per_pass_and_preserves_all_xml(self):
        wb = representative_workbook()
        # Reintroducing openpyxl membership scans fails deterministically, instead
        # of relying on a wall-clock threshold affected by the test machine.
        with patch.object(MergedCellRange, "__contains__", side_effect=AssertionError("Repeated range parsing")), \
                patch.object(paper, "_merged_cell_extents", wraps=paper._merged_cell_extents) as build:
            paper.apply_plain_presentation(wb.active)
            self.assertEqual(build.call_count, 1)
            # openpyxl populates column outline metadata during its first save.
            workbook_parts(wb)
            before = workbook_parts(wb)
            paper.validate_plain_workbook(wb)
            self.assertEqual(build.call_count, 2)
            self.assertEqual({name: hashlib.sha256(data).hexdigest() for name, data in before.items()},
                             {name: hashlib.sha256(data).hexdigest() for name, data in workbook_parts(wb).items()})
        self.assertEqual(wb.active["C3"].value, "=SUM(D3:E3)")

    def test_all_coordinates_match_range_scan_including_merged_interiors(self):
        wb = representative_workbook()
        ws = wb.active
        index = paper._merged_cell_extents(ws)
        for row in ws:
            for cell in row:
                expected = (cell.column, cell.column, cell.row, cell.row)
                for merged in ws.merged_cells.ranges:
                    if merged.min_row <= cell.row <= merged.max_row and merged.min_col <= cell.column <= merged.max_col:
                        expected = (merged.min_col, merged.max_col, merged.min_row, merged.max_row)
                        break
                self.assertEqual(paper._cell_extent(ws, cell, index), expected)
        ws.unmerge_cells("A1:B1")
        self.assertEqual(paper._cell_extent(ws, ws["A1"]), (1, 1, 1, 1))
        ws.merge_cells("A1:C2")
        self.assertEqual(paper._cell_extent(ws, ws["A1"]), (1, 3, 1, 2))

    def test_invalid_presentation_still_fails(self):
        def grid(ws): ws.sheet_view.showGridLines = False
        def fit(ws): ws.sheet_properties.pageSetUpPr.fitToPage = True
        def width(ws): ws.page_setup.fitToWidth = 1
        def height(ws): ws.page_setup.fitToHeight = 1
        def area(ws): ws.print_area = "A1:B2"
        def fill(ws): ws["C3"].fill = PatternFill("solid", fgColor="FFFFFF")
        def border(ws): ws["C3"].border = Border(bottom=Side(style="thin"))
        def color(ws): ws["C3"].font = Font(color="FF0000")
        def wrap(ws): ws["C3"].alignment = Alignment(wrap_text=False)
        def clipped(ws): ws.row_dimensions[3].height = 1
        for mutation in (grid, fit, width, height, area, fill, border, color, wrap, clipped):
            with self.subTest(rule=mutation.__name__):
                wb = Workbook()
                ws = wb.active
                ws.merge_cells("A1:B1")
                ws["A1"], ws["C3"] = "Header", 0.1234
                paper.apply_plain_presentation(ws)
                mutation(ws)
                with self.assertRaises(ValueError):
                    paper.validate_plain_workbook(wb)


if __name__ == "__main__":
    unittest.main()
