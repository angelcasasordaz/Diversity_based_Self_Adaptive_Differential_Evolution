"""Validated cache-only reporting and append-only report version publication."""
import argparse
from contextlib import contextmanager
from dataclasses import dataclass
import hashlib
import json
import os
from pathlib import Path
import pickle
import re
import shutil
import sys
import tempfile
import time

import numpy as np


def framework():
    # main_best.py may be running as __main__; reuse that exact module.
    main = sys.modules.get("__main__")
    if getattr(main, "__file__", "").endswith("main_best.py"):
        return main
    import main_best
    return main_best


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


@contextmanager
def report_stage(label):
    """Make long cache-only operations visible even in a buffered IDE console."""
    started = time.perf_counter()
    print(f'[report] {label} ...', flush=True)
    try:
        yield
    except BaseException:
        print(f'[report] {label} failed after {time.perf_counter() - started:.1f}s', flush=True)
        raise
    else:
        print(f'[report] {label} complete ({time.perf_counter() - started:.1f}s)', flush=True)


@contextmanager
def report_guard(destinations):
    """Reject scientific execution and filesystem mutations outside replica folders."""
    state = {"active": True, "optimization_calls": 0}
    forbidden = {"_run_single", "execute_pending_runs", "build_optimizer",
                 "configure_compute_backend", "save_cache", "start_gpu_request_service",
                 "load_dataset", "load_codesmell_dataset", "load_mafese_dataset"}
    blocked_codes = {}
    def profile(frame, event, arg):
        if event != 'call':
            return
        # Code metadata is immutable. Inspect it once instead of re-triggering
        # object.__getattr__ audit events on every library function call.
        code = frame.f_code
        blocked = blocked_codes.get(code)
        if blocked is None:
            name = code.co_name
            blocked = (name in forbidden or (name == 'solve' and 'mealpy' in code.co_filename)
                       or (name == 'fit' and 'mafese' in code.co_filename))
            blocked_codes[code] = blocked
        if blocked:
            state["optimization_calls"] += 1
            raise RuntimeError(f"REPORT-ONLY blocked scientific execution: {code.co_name}")
    def allowed(path):
        if isinstance(path, int):
            return False
        resolved = Path(os.fsdecode(path)).resolve()
        return any(resolved == base or base in resolved.parents for base in destinations) and resolved.suffix.lower() not in {'.pkl', '.pickle', '.ckpt', '.pdf'}
    mutation_events = frozenset({'open', 'os.remove', 'os.rmdir', 'os.mkdir', 'os.chmod',
                                'os.utime', 'os.truncate', 'shutil.rmtree', 'os.rename',
                                'os.link', 'os.symlink', 'subprocess.Popen', 'os.system', 'os.fork'})
    def audit(event, args):
        if not state["active"] or event not in mutation_events:
            return
        targets = []
        def at(path, dir_fd):
            # shutil.rmtree uses descriptor-relative unlink/rmdir on Linux.
            if dir_fd is not None and dir_fd != -1 and not os.path.isabs(path):
                return Path(os.readlink(f"/proc/self/fd/{dir_fd}")) / os.fsdecode(path)
            return path
        if event == "open":
            path, mode, flags = args
            if (flags & (os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_TRUNC | os.O_APPEND)):
                targets = [path]
        elif event in {"os.remove", "os.rmdir"}:
            targets = [at(args[0], args[1])]
        elif event in {"os.mkdir", "os.chmod"}:
            targets = [at(args[0], args[2])]
        elif event in {"os.utime", "os.truncate", "shutil.rmtree"}:
            targets = [args[0]]
        elif event in {"os.rename", "os.link", "os.symlink"}:
            targets = list(args[:2])
        elif event in {"subprocess.Popen", "os.system", "os.fork"}:
            raise RuntimeError(f"REPORT-ONLY blocked process execution: {event}")
        if any(not allowed(path) for path in targets):
            raise PermissionError(f"REPORT-ONLY blocked mutation outside replica output: {event} {targets}")
    previous = sys.getprofile()
    previous_bytecode = sys.dont_write_bytecode
    sys.dont_write_bytecode = True
    sys.addaudithook(audit)
    sys.setprofile(profile)
    try:
        yield state
    finally:
        state["active"] = False
        sys.setprofile(previous)
        sys.dont_write_bytecode = previous_bytecode


@dataclass
class CompletedReport:
    args: argparse.Namespace
    results: dict
    indexed: dict
    datasets: list
    classifiers: list
    algorithms: list
    metrics: list
    signature: str
    sources: dict

    @property
    def exp_tag(self):
        return f"EXP{self.args.exp_id:03d}"


def safe_path(path):
    """Reject redirects, traversal and non-directory ancestors before any write."""
    path = Path(path).absolute()
    if '..' in path.parts or any(p.is_symlink() for p in (path, *path.parents)):
        raise ValueError(f"Unsafe reporting path: {path}")
    if any(p.exists() and not p.is_dir() for p in path.parents):
        raise ValueError(f"Non-directory reporting ancestor: {path}")
    return path


def load_completed_cache(args):
    """Read exact final-cache identities; never use progress or migration readers.

    Configuration determines the required grid, and cache labels/records supply
    its data. All framework run fields are required by the numerical exporters.
    Unavailable metrics/curves are explicitly omitted; partial metric grids fail.
    """
    if args.experiment_mode == 'full' and getattr(args, 'report_current_science', False):
        return load_cached_figure_report(args)
    from reporting.paper_tables import METRICS
    m = framework()
    args = argparse.Namespace(**vars(args))
    if not isinstance(args.exp_id, int) or args.exp_id < 0 or args.runs < 1:
        raise ValueError("An explicit nonnegative EXP ID and positive run count are required")
    m.apply_experiment_mode(args)  # Local Namespace only; no scientific globals change.
    signature = m.build_cache_signature(args)
    root = safe_path(args.output_root)
    tag = f"EXP{args.exp_id:03d}"
    cache = safe_path(root / 'Results' / tag / args.experiment_mode / 'cache')
    if not cache.is_dir():
        raise FileNotFoundError(f"Missing completed cache: {cache}; no recomputation allowed")
    configured = m.configured_dataset_names(args)
    classifiers = list(dict.fromkeys(str(c).lower() for c in args.estimators))
    if not classifiers:
        raise ValueError("No configured classifiers")
    if configured is None:
        # Resolve the configured source list without reading dataset contents.
        # This detects a whole missing dataset, as well as missing classifiers.
        configured = [spec.name for spec in m.resolve_dataset_specs(args)]
    datasets = list(dict.fromkeys(Path(str(d)).stem if str(d).endswith('.csv') else str(d)
                                  for d in configured))
    if not datasets:
        raise ValueError("No configured datasets")
    if any('/' in d or '\\' in d or d in {'.', '..'} for d in datasets):
        raise ValueError("Unsafe dataset identity")
    results, indexed, sources, algorithms = {}, {}, {}, []
    show_tf, show_cls = len(args.transfer_functions) > 1, len(classifiers) > 1
    for dataset in datasets:
        results[dataset] = {}
        for classifier in classifiers:
            legacy_full = False
            classifier_signature = m.classifier_cache_signature(args, signature, classifier)
            path = safe_path(cache / f'{tag}_{dataset}_{classifier}_{classifier_signature}_results.pkl')
            if not path.is_file() and classifier == 'rf':
                path = safe_path(cache / f'{tag}_{dataset}_{classifier}_{signature}_results.pkl')
            if not path.is_file() and args.experiment_mode == 'full':
                # Historical reporting stays exact and read-only; wildcard
                # migration belongs exclusively to normal scientific execution.
                legacy_signature = m.build_legacy_cache_signature(args)
                path = safe_path(cache / f'{tag}_{dataset}_{classifier}_{legacy_signature}_results.pkl')
                legacy_full = True
            if not path.is_file():
                raise FileNotFoundError(f"Missing completed cache: {path}; no recomputation allowed")
            before = sha256(path)
            payload = m.load_cache(str(path))
            expected = m.expected_result_labels(args, classifier, show_tf, show_cls)
            canonical = m.expected_result_labels(args, classifier, show_tf, True)
            groups = m.expected_result_labels(args, classifier, show_tf, False)
            if args.experiment_mode == 'full' and not legacy_full:
                # FULL updates append snapshots. Keep reporting read-only and
                # final-only, selecting a complete scientifically matching grid.
                snapshot_signature = classifier_signature if path.name.startswith(f'{tag}_{dataset}_{classifier}_{classifier_signature}_') else signature
                finals = sorted(cache.glob(f'{tag}_{dataset}_{classifier}_{snapshot_signature}_final_*_results.pkl'))
                for candidate in (path, *finals):
                    candidate = safe_path(candidate)
                    stored = m.load_cache(str(candidate))
                    selected = {}
                    for (label, legacy), (output, output_legacy), (group, group_legacy) in zip(expected, canonical, groups):
                        keys = [key for key in dict.fromkeys((label, legacy, output, output_legacy, group, group_legacy))
                                if isinstance(stored, dict) and key in stored]
                        if len(keys) != 1:
                            break
                        row = stored[keys[0]]
                        parsed = m.parse_result_label(output, args)
                        metadata = m.full_checkpoint_metadata(args, parsed['method'], dataset, classifier,
                            parsed['transfer_function'] or args.transfer_functions[0])
                        values, reason = m.validate_source_label_runs(row, classifier, args.runs, metadata)
                        if reason or len(values['AccRuns']) != args.runs:
                            break
                        selected[keys[0]] = row
                    else:
                        path, payload, before = candidate, selected, sha256(candidate)
                        break
            if not isinstance(payload, dict) or len(payload) != len(expected):
                raise ValueError(f"Incomplete/extra optimizer rows: {path}")
            used = set()
            for (label, legacy), (output_label, output_legacy), (algorithm, group_legacy) in zip(expected, canonical, groups):
                # Classifier-specific legacy caches can have unqualified labels.
                candidates = (label, legacy, output_label, output_legacy, algorithm, group_legacy)
                matches = [key for key in dict.fromkeys(candidates) if key in payload]
                if len(matches) != 1 or matches[0] in used:
                    raise ValueError(f"Missing/ambiguous result label: {path}/{label}")
                used.add(matches[0])
                row = payload[matches[0]]
                if not isinstance(row, dict):
                    raise ValueError(f"Invalid cache row: {path}/{label}")
                if classifier == 'rf':
                    backend_reason = m.rf_backend_compatibility(row, m.resolve_rf_backend(args))
                    if backend_reason:
                        raise ValueError(f'{path.name}/{label}: {backend_reason}')
                fields = [metric.run_key for metric in METRICS] + ['CurvesAll']
                if all(field in row for field in fields):
                    values, reason = m.validate_source_label_runs(row, classifier, args.runs)
                else:
                    # The framework validator requires its full current schema.
                    # Older caches may omit an entire metric: validate only actual
                    # observations, and never fabricate fields to pass that validator.
                    values = {field: row[field] for field in fields if field in row}
                    reason = None if str(row.get('Estimator', '')).lower() == classifier else 'classifier metadata mismatch'
                if reason or int(row.get('CompletedRuns', 0)) != args.runs:
                    raise ValueError(f"Incomplete cache {path}/{label}: {reason or 'completed run count mismatch'}")
                parsed = m.parse_result_label(output_label, args)
                if args.experiment_mode == 'full' and (not legacy_full or 'CacheIdentity' in row):
                    expected_metadata = m.full_checkpoint_metadata(
                        args, parsed['method'], dataset, classifier,
                        parsed['transfer_function'] or args.transfer_functions[0],
                    )
                    if not m.checkpoint_metadata_matches(row, expected_metadata):
                        raise ValueError(f'Scientific cache identity mismatch: {path}/{label}')
                if parsed['estimator'] != classifier:
                    raise ValueError(f"Classifier identity mismatch: {output_label}")
                if row.get('ExperimentMode', args.experiment_mode) != args.experiment_mode:
                    raise ValueError(f"Mode identity mismatch: {output_label}")
                for metric in METRICS:
                    if metric.run_key not in values:
                        continue
                    arr = np.asarray(values[metric.run_key], dtype=float)
                    if arr.shape != (args.runs,) or not np.isfinite(arr).all():
                        raise ValueError(f"Invalid run values: {path}/{label}/{metric.run_key}")
                    mean_key = metric.run_key.replace('Runs', 'Mean')
                    if not np.isclose(row.get(mean_key, np.nan), arr.mean(), rtol=1e-12, atol=1e-12):
                        raise ValueError(f"Cached summary disagrees with runs: {path}/{label}/{mean_key}")
                curves = [np.asarray(c, dtype=float) for c in values.get('CurvesAll', [])]
                if 'CurvesAll' in values and len(curves) != args.runs:
                    raise ValueError(f'Incomplete convergence run array: {label}')
                if any(c.ndim != 1 or not np.isfinite(c).all() for c in curves):
                    raise ValueError(f"Invalid convergence curves: {label}")
                mean_curve = np.asarray(row.get('Curve', []), dtype=float)
                if mean_curve.size:
                    if (mean_curve.shape != (args.epochs,) or not np.isfinite(mean_curve).all()
                            or (curves and (not all(c.size for c in curves) or not np.allclose(
                                mean_curve, m.pad_mean_curves(curves, args.epochs), rtol=1e-12, atol=1e-12)))):
                        raise ValueError(f"Invalid cached mean convergence: {label}")
                key = dataset, classifier, algorithm
                if key in indexed:
                    raise ValueError(f"Ambiguous scientific identity: {key}")
                indexed[key] = row
                results[dataset][output_label] = row
                if algorithm not in algorithms:
                    algorithms.append(algorithm)
            if set(payload) != used or sha256(path) != before:
                raise ValueError(f"Cache identity/content changed during read: {path}")
            sources[str(path.relative_to(root))] = before
    metrics = [metric for metric in METRICS if any(metric.run_key in row for row in indexed.values())]
    if not metrics:
        raise ValueError('No available run metrics in completed cache')
    for metric in metrics:
        if any(metric.run_key not in row for row in indexed.values()):
            raise ValueError(f'Incomplete metric grid: {metric.run_key}')
    return CompletedReport(args, results, indexed, datasets, classifiers, algorithms,
                           metrics, signature, sources)


def report_versions(root, exp_id):
    """Inspect both trees. Unpaired, redirected or malformed versions are errors."""
    if not isinstance(exp_id, int) or exp_id < 0:
        raise ValueError('A nonnegative user-selected EXP ID is required')
    root = safe_path(root)
    versions = []
    for kind in ('Figures', 'Results'):
        parent = safe_path(root / kind / f'EXP{exp_id:03d}')
        found = set()
        if parent.exists():
            for path in parent.iterdir():
                match = re.fullmatch(r'full_rep([1-9][0-9]*)', path.name)
                if not match:
                    if re.fullmatch(r'full_rep[0-9]+', path.name):
                        raise ValueError(f"Noncanonical report version: {path}")
                    continue
                safe_path(path)
                if not path.is_dir():
                    raise ValueError(f"Conflicting report version: {path}")
                for child in path.rglob('*'):
                    if child.is_symlink():
                        raise ValueError(f"Symlink within report: {child}")
                manifest_path = path / 'validation.json'
                if manifest_path.is_file():
                    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
                    if (manifest.get('experiment_id', exp_id) != exp_id
                            or manifest.get('report_version', int(match[1])) != int(match[1])):
                        raise ValueError(f'Conflicting report manifest identity: {manifest_path}')
                found.add(int(match[1]))
        versions.append(found)
    if versions[0] != versions[1]:
        raise ValueError(f"Incomplete Figures/Results report pairs: {versions}; existing files preserved")
    return sorted(versions[0])


def next_report_version(root, exp_id):
    return max(report_versions(root, exp_id), default=0) + 1


def _rename_new(source, destination):
    """Atomic, no-replace directory publication (Linux); fail closed elsewhere."""
    import ctypes
    libc = ctypes.CDLL(None, use_errno=True)
    rename = getattr(libc, 'renameat2', None)
    if rename is None:
        raise RuntimeError('Atomic no-replace publication requires renameat2')
    if rename(-100, os.fsencode(source), -100, os.fsencode(destination), 1):
        error = ctypes.get_errno()
        raise OSError(error, os.strerror(error), str(destination))


@contextmanager
def staged_version(root, exp_id):
    """Serialize allocators, stage both trees, then publish a matching pair.

    A crash between the two renames leaves an unpaired version: the next attempt
    fails safely for manual inspection. Temporary names never consume a number.
    """
    root = safe_path(root)
    next_report_version(root, exp_id)  # Preflight before directory creation.
    parents = [safe_path(root / k / f'EXP{exp_id:03d}') for k in ('Figures', 'Results')]
    for parent in parents:
        parent.mkdir(parents=True, exist_ok=True)
    lock = parents[1] / '.report-allocation.lock'
    try:
        lock.mkdir()
    except FileExistsError as exc:
        raise ValueError(f'Report allocation is locked; inspect before retrying: {lock}') from exc
    stages, published = [], []
    try:
        version = next_report_version(root, exp_id)
        finals = [p / f'full_rep{version}' for p in parents]
        for parent in parents:
            stages.append(Path(tempfile.mkdtemp(prefix='.report-staging-', dir=parent)))
        yield version, stages[0], stages[1], finals
        for source, target in zip(stages, finals):
            safe_path(target)
            inode = source.stat().st_ino
            _rename_new(source, target)
            published.append((source, target, inode))
    except BaseException:
        for source, target, inode in reversed(published):
            if not target.is_symlink() and target.stat().st_ino == inode:
                _rename_new(target, source)
        raise
    finally:
        for stage in stages:
            if stage.exists() and not stage.is_symlink():
                shutil.rmtree(stage)
        lock.rmdir()


def _validate_artifacts(figures, results, required):
    from openpyxl import load_workbook
    from PIL import Image
    for relative in required:
        if not (results / relative).is_file():
            raise ValueError(f'Missing required report output: {relative}')
    hashes = {}
    for base in (figures, results):
        for path in base.rglob('*'):
            if path.is_symlink():
                raise ValueError(f'Redirected generated output: {path}')
            if not path.is_file():
                continue
            if path.suffix.lower() in {'.pdf', '.pkl', '.pickle', '.ckpt'} or path.stat().st_size == 0:
                raise ValueError(f'Invalid generated output: {path}')
            if path.suffix == '.xlsx':
                wb = load_workbook(path, data_only=False)
                try:
                    if not wb.sheetnames or any(ws.max_row < 1 for ws in wb):
                        raise ValueError(f'Empty workbook: {path}')
                    if path.name.startswith('Paper_Tables'):
                        from reporting.paper_tables import validate_plain_workbook
                        validate_plain_workbook(wb)
                finally:
                    wb.close()
            if path.suffix == '.png':
                with Image.open(path) as png:
                    if png.format != 'PNG' or any(abs(v - 600) > .1 for v in png.info.get('dpi', (0, 0))):
                        raise ValueError(f'Expected 600 dpi PNG: {path}')
                    png.verify()
            hashes[('Figures/' if base == figures else 'Results/') + str(path.relative_to(base))] = sha256(path)
    return hashes


def generate_outputs(report, figures, results):
    """Reuse framework Excel exporters without calling its experiment dispatcher."""
    from reporting import figures as plotting, statistics, paper_tables
    m, args = framework(), report.args
    figures.mkdir(parents=True, exist_ok=True)
    results.mkdir(parents=True, exist_ok=True)
    tag = report.exp_tag
    required = [f'Global_Results_{tag}.xlsx', f'Statistical_Results_{tag}.xlsx', f'Paper_Tables_{tag}.xlsx',
                'statistics/statistical_summary.csv', 'statistics/pairwise_wilcoxon_holm.csv',
                'statistics/statistical_report.txt']
    with report_stage(f'{tag}: Global Results Excel'):
        m.export_global_excel(report.results, report.datasets, str(results / required[0]))
    with report_stage(f'{tag}: Statistical Results Excel'):
        m.export_statistical_excel(report.results, report.datasets, args.optimizers, args, str(results / required[1]))
    with report_stage(f'{tag}: paper tables (format, serialize, validate)'):
        paper_tables.export_indexed_tables(report.indexed, report.datasets, report.algorithms, report.classifiers,
                                          report.metrics, results / required[2], title=f'{tag} {args.experiment_mode.upper()}')
    skipped = [{'output': f'{metric.name} tables/figures', 'reason': f'{metric.run_key} unavailable in every cache row'}
               for metric in paper_tables.METRICS if metric not in report.metrics]
    if args.experiment_mode in {'full', 'ablation'} and any(metric.run_key == 'FitRuns' for metric in report.metrics):
        name = f'{args.experiment_mode.capitalize()}_Friedman_Analysis_{tag}.xlsx'
        with report_stage(f'{tag}: Friedman Excel'):
            m.export_friedman_analysis(report.results, report.datasets, args.optimizers, args, str(results / name))
        required.append(name)
    else:
        skipped.append({'output': 'Framework fitness Friedman Excel', 'reason': 'Requires full/ablation mode and available FitRuns'})
    with report_stage(f'{tag}: publication figures at 600 dpi'):
        skipped.extend(plotting.generate(report, figures))
    with report_stage(f'{tag}: statistical analysis and figures'):
        skipped.extend(statistics.export(report, plotting.statistics_destination(report, figures), results / 'statistics'))
    if args.experiment_mode == 'full':
        actual = {p.name for p in figures.iterdir() if p.is_file()}
        expected = {name for name in plotting.base_figure_names(report)
                    if plotting.full_figure_path(figures, name).parent == figures}
        omitted = {item['output'] for item in skipped} & expected
        if actual != expected - omitted:
            raise ValueError(f'Unexpected root figure layout: expected {sorted(expected - omitted)}, got {sorted(actual)}')
    return required, skipped


def run_report(args):
    """Create one complete report version for the user-selected EXP and modes."""
    if getattr(args, 'report_only', False) or getattr(args, 'figure_language', framework().FIGURE_LANGUAGE) == 'es':
        return run_figure_report(args)
    # Import exporters/dependencies before the strict mutation guard is active.
    started = time.perf_counter()
    print(f'[report] Starting cache-only EXP{args.exp_id:03d}: {", ".join(args.experiment_modes)}', flush=True)
    from reporting import figures, statistics, paper_tables
    m = framework()
    root = safe_path(args.output_root)
    modes = list(dict.fromkeys(args.experiment_modes))
    if not modes or any(mode not in {'full', 'ablation', 'sensitivity', 'sensitivity_weights', 'transfer_functions'} for mode in modes):
        raise ValueError('Select at least one supported experiment mode')
    reports = []
    with report_stage('Validate original completed caches'), report_guard(()) as preflight:
        for mode in modes:
            local = m.clone_args_for_mode(args, mode)
            studies = m.sensitivity_study_args(local) if mode == 'sensitivity' else [local]
            from reporting.experiment_config import read_manifest
            for study in studies:
                if read_manifest(study) is not None:
                    m.apply_experiment_mode(study)
                    reports.append(load_cached_figure_report(study))
                else:
                    reports.append(load_completed_cache(study))
    with report_stage('Hash protected historical outputs'):
        previous = {str(p.relative_to(root)): sha256(p) for kind in ('Figures', 'Results')
                    for p in (root / kind / f'EXP{args.exp_id:03d}').rglob('*') if p.is_file()}
    destination_root = safe_path(getattr(args, 'report_output_root', None) or root)
    print(f'[report] Source: {root}; destination: {destination_root}', flush=True)
    with staged_version(destination_root, args.exp_id) as (version, fig, res, finals):
        old_tempdir = tempfile.tempdir
        tempfile.tempdir = str(res)
        try:
            with report_guard((fig, res)) as guard:
                required, entries = [], []
                for report in reports:
                    sub = Path('.') if len(reports) == 1 else Path(report.args.experiment_mode)
                    if len(reports) > 1 and report.args.experiment_mode == 'sensitivity':
                        sub /= report.args.sensitivity_parameter
                    names, skipped = generate_outputs(report, fig / sub, res / sub)
                    required.extend(str(sub / name) for name in names)
                    entries.append({'mode': report.args.experiment_mode, 'subdirectory': str(sub),
                                    'datasets': report.datasets, 'classifiers': report.classifiers,
                                    'algorithms': report.algorithms, 'metrics': [m.run_key for m in report.metrics],
                                    'source_directory': str(root / 'Results' / report.exp_tag / report.args.experiment_mode / 'cache'),
                                    'completed_runs': sorted({r['CompletedRuns'] for r in report.indexed.values()}),
                                    'cache_identity': report.signature, 'source_cache_sha256': report.sources,
                                    'figure_style': figures.base_style.STYLE_ID,
                                    'base_classifier': figures.base_classifier(report) if report.args.experiment_mode == 'full' else None,
                                    'base_metric': figures.base_metric_token(report) if report.args.experiment_mode == 'full' else None,
                                    'figure_language': figures.figure_language(report) if report.args.experiment_mode == 'full' else None,
                                    'root_figures': sorted(p.name for p in (fig / sub).glob('*.png')),
                                    'individual_figures': sorted(str(p.relative_to(fig / sub)) for p in
                                        figures.individual_destination(report, fig / sub).rglob('*.png'))
                                        if report.args.experiment_mode == 'full' else [],
                                    'statistical_figures': sorted(str(p.relative_to(fig / sub)) for p in
                                        figures.statistics_destination(report, fig / sub).rglob('*.png')),
                                    'skipped_outputs': skipped})
                with report_stage('Validate every generated artifact'):
                    hashes = _validate_artifacts(fig, res, required)
                with report_stage('Verify protected historical hashes'):
                    if any(sha256(root / path) != digest for path, digest in previous.items()):
                        raise ValueError('Protected source/report changed during reporting')
                manifest = {'experiment_id': args.exp_id, 'report_version': version,
                            'report_version_is_scientific_repetition': False, 'reports': entries,
                            'figures_destination': str(finals[0]), 'results_destination': str(finals[1]),
                            'outputs_sha256': hashes, 'required_results': required,
                            'optimization_calls': guard['optimization_calls'] + preflight['optimization_calls'],
                            'dpi': 600, 'pdfs_generated': 0, 'protected_files_unchanged': True}
                (res / 'validation.json').write_text(json.dumps(manifest, indent=2), encoding='utf-8')
        finally:
            tempfile.tempdir = old_tempdir
    print(f'[report] Published EXP{args.exp_id:03d} full_rep{version} in {time.perf_counter() - started:.1f}s: '
          f'{finals[0]} and {finals[1]}', flush=True)
    from reporting.experiment_config import save_report_config
    for report in reports:
        save_report_config(report)
    return manifest


REPORT_SCIENCE_FIELDS = ('runs', 'epochs', 'pop_size', 'test_size', 'random_state', 'seed_base')


def _audit_cache(status, field, requested, stored, evidence=''):
    print(f'[cache] {status} {field}: requested={requested!r}; stored={stored!r}'
          + (f'; {evidence}' if evidence else ''), flush=True)


def _legacy_report_identity(args, signature, methods):
    """Verify recorded historical digests, never guess unrecorded numeric settings.

    Pre-revision FULL comparisons used the same numeric fields, but included
    global DSA/sensitivity settings even for other methods (bcc0cb9 schema).
    Test only the current and documented historical auxiliary defaults. Neither
    a matching dataset nor a completed row alone proves pop/split/seed settings.
    """
    m = framework()
    historical = argparse.Namespace(**vars(args))
    historical.optimizers = methods
    payload = {key: getattr(args, key) for key in REPORT_SCIENCE_FIELDS}
    payload.update(experiment_mode=args.experiment_mode, optimizers=methods,
                   transfer_functions=list(args.transfer_functions), obj_name='AS',
                   fitness_mode='minimize_metric_loss_plus_feature_ratio_v1')
    auxiliary = {key: getattr(args, key) for key in ('dsade_beta_min', 'dsade_beta_max',
        'dsade_pcr', 'dsade_mahal_q', 'sensitivity_parameter', 'sensitivity_values')}
    contracts = [auxiliary]
    for low, high in ((.10, .60), (.40, .80)):
        contracts.append(dict(dsade_beta_min=low, dsade_beta_max=high, dsade_pcr=.10,
                              dsade_mahal_q=.50, sensitivity_parameter='mahalanobis_q',
                              sensitivity_values=[.50, .68, .80, .90]))
    for contract in contracts:
        for key, value in contract.items():
            setattr(historical, key, value)
        # The newer legacy schema includes revisioned optimizer identities.
        if m.build_legacy_cache_signature(historical) == signature:
            args.report_legacy_parameters = contract
            return 'verified legacy digest (revisioned schema)'
        older = {**payload, **contract}
        if args.experiment_mode == 'sensitivity_weights':
            older.update(fitness_alpha=args.fitness_alpha, fitness_beta=args.fitness_beta,
                         sensitivity_weight_pairs=[list(pair) for pair in args.sensitivity_weight_pairs])
        digest = hashlib.sha1(json.dumps(older, sort_keys=True).encode()).hexdigest()[:10]
        if args.experiment_mode == 'sensitivity':
            digest = f'{args.sensitivity_parameter}_{digest}'
        if digest == signature:
            args.report_legacy_parameters = contract
            return 'verified legacy digest (pre-revision schema)'
    raise ValueError('CacheIdentity absent; legacy digest cannot verify '
                     'runs/epochs/pop_size/test_size/random_state/seed_base for the requested configuration')


def _report_algorithm(label, classifier):
    suffix = f'_{classifier.upper()}'
    group = label[:-len(suffix)] if classifier and label.upper().endswith(suffix) else label
    method = re.split(r'_(?:[VS]STF_|SENS_|WEIGHTS_A)', group, maxsplit=1, flags=re.I)[0]
    canonical = framework().resolve_optimizer_name(method)
    normalized = framework().optimizer_acronym(canonical).upper() + group[len(method):].upper()
    return normalized, canonical


def _validate_report_row(row, args, dataset, classifier, method, legacy_evidence):
    from reporting.paper_tables import METRICS
    if str(row.get('Estimator', '')).lower() != classifier:
        raise ValueError('estimators mismatch in row metadata')
    if row.get('ExperimentMode', args.experiment_mode) != args.experiment_mode:
        raise ValueError('experiment_mode mismatch in row metadata')
    identity = row.get('CacheIdentity')
    if identity is not None:
        if not isinstance(identity, dict):
            raise ValueError('Invalid CacheIdentity')
        expected = {key: getattr(args, key) for key in REPORT_SCIENCE_FIELDS + ('fitness_alpha', 'fitness_beta')}
        expected.update(experiment_mode=args.experiment_mode, dataset_name=dataset,
                        classifier=classifier, optimizer=method)
        if '--dataset-source' in getattr(args, 'report_explicit_options', ()):
            expected['dataset_source'] = args.dataset_source
        for key, value in expected.items():
            actual = identity.get(key)
            if actual != value:
                raise ValueError(f'{key} MISMATCH: requested={value!r}, stored={actual!r} (CacheIdentity)')
        tf = identity.get('transfer_function')
        if tf not in args.transfer_functions:
            raise ValueError(f'transfer_functions MISMATCH: requested={args.transfer_functions}, stored={tf!r}')
        if getattr(args, 'report_current_science', False):
            expected_identity = framework().full_cache_identity(args, method, dataset, classifier, tf)
            if identity != expected_identity:
                raise ValueError(f'optimizer scientific identity MISMATCH for {method}')
        parameters = identity.get('optimizer_parameters', {})
        manifest = getattr(args, 'report_manifest', None)
        if manifest is not None:
            stored_parameters = manifest.get('optimizer_parameters', {}).get(method)
            if stored_parameters is not None and stored_parameters != parameters:
                raise ValueError(f'optimizer_parameters MISMATCH for {method} against experiment_config.json')
        for key, field in (('epoch', 'epochs'), ('pop_size', 'pop_size')):
            if key in parameters and parameters[key] != getattr(args, field):
                raise ValueError(f'{field} MISMATCH in optimizer_parameters: {parameters[key]!r}')
    elif not legacy_evidence:
        raise ValueError('CacheIdentity absent and no verified legacy identity')
    elif args.fitness_alpha != .9 or args.fitness_beta != .1:
        raise ValueError('fitness_alpha/fitness_beta MISMATCH: legacy cache requires the historical 0.9/0.1 contract')
    runs = int(row.get('CompletedRuns', 0))
    if runs != args.runs or row.get('CompletedRuns') != runs:
        raise ValueError(f'runs MISMATCH: requested={args.runs}, completed={runs}')
    for metric in METRICS:
        if metric.run_key not in row:
            continue
        values = np.asarray(row[metric.run_key], dtype=float)
        if values.shape != (runs,) or not np.isfinite(values).all():
            raise ValueError(f'Invalid cached runs: {metric.run_key}')
        mean = row.get(metric.run_key.replace('Runs', 'Mean'), np.nan)
        if not np.isclose(mean, values.mean(), rtol=1e-12, atol=1e-12):
            raise ValueError(f'Cached summary disagrees with runs: {metric.run_key}')
    curve = np.asarray(row.get('Curve', []), dtype=float)
    if curve.ndim != 1 or not np.isfinite(curve).all():
        raise ValueError('Invalid cached curve')
    if curve.size and curve.size != args.epochs:
        raise ValueError(f'epochs MISMATCH: requested={args.epochs}, curve_length={curve.size}')
    if 'CurvesAll' in row:
        curves = [np.asarray(c, dtype=float) for c in row['CurvesAll']]
        if len(curves) != runs or any(c.ndim != 1 or len(c) > args.epochs or not np.isfinite(c).all()
                                    for c in curves):
            raise ValueError('Incomplete or invalid cached convergence runs')


def _same_report_observations(left, right):
    from reporting.paper_tables import METRICS
    for field in [m.run_key for m in METRICS] + ['Curve']:
        if (field in left) != (field in right) or not np.array_equal(left.get(field, []), right.get(field, [])):
            return False
    a, b = left.get('CurvesAll', []), right.get('CurvesAll', [])
    return (left.get('CacheIdentity') == right.get('CacheIdentity') and len(a) == len(b)
            and all(np.array_equal(x, y) for x, y in zip(a, b)))


def load_cached_figure_report(args, *, use_manifest=True):
    """Discover compatible completed caches of any hash, including progress files."""
    from reporting.paper_tables import METRICS
    m = framework()
    local = argparse.Namespace(**vars(args))
    if local.experiment_mode == 'full' and getattr(local, 'report_current_science', False):
        paths = m.make_read_only_source_paths(local, include_current=True)
        datasets = [spec.name for spec in m.resolve_dataset_specs(local)]
        classifiers = list(local.estimators)
        results, indexed, sources, algorithms, missing = {}, {}, {}, [], []
        for dataset in datasets:
            results[dataset] = {}
            for classifier in classifiers:
                payload, _, origins = m.load_compatible_full_cache_payload(
                    paths, local, dataset, classifier, final_only=True, return_origins=True)
                expected = m.expected_result_labels(local, classifier,
                    len(local.transfer_functions) > 1, len(classifiers) > 1)
                for label, _ in expected:
                    row = (payload or {}).get(label)
                    if row is None:
                        missing.append(f'{dataset}/{classifier}/{label}')
                        continue
                    algorithm, _ = _report_algorithm(label, classifier)
                    results[dataset][label] = row
                    indexed[dataset, classifier, algorithm] = row
                    if algorithm not in algorithms:
                        algorithms.append(algorithm)
                    path = safe_path(origins[label])
                    sources[str(path.relative_to(Path(local.output_root).absolute()))] = sha256(path)
        if missing:
            raise ValueError('Cache MISMATCH: missing compatible complete final rows: ' + ', '.join(missing))
        metrics = [metric for metric in METRICS if all(metric.run_key in row for row in indexed.values())]
        return CompletedReport(local, results, indexed, datasets, classifiers, algorithms, metrics,
                               'optimizer-local-final-caches', sources)
    from reporting.experiment_config import read_manifest, apply_manifest, manifest_path
    config = read_manifest(local) if use_manifest else None
    if config is not None:
        local = apply_manifest(local, config)
    root = safe_path(local.output_root)
    tag = f'EXP{local.exp_id:03d}'
    cache = safe_path(root / 'Results' / tag / local.experiment_mode / 'cache')
    if not cache.is_dir():
        raise FileNotFoundError(f'Missing completed cache: {cache}; no recomputation allowed')
    if local.runs < 1 or local.epochs < 1:
        raise ValueError('Positive runs and epochs are required')
    explicit = getattr(local, 'report_explicit_options', ())
    requested_datasets = [Path(d).stem for d in local.datasets] if local.datasets is not None else None
    if requested_datasets is None and '--dataset-source' in explicit:
        requested_datasets = m.configured_dataset_names(local)
        if requested_datasets is None:
            requested_datasets = [spec.name for spec in m.resolve_dataset_specs(local)]
    requested_classifiers = [str(c).lower() for c in local.estimators] if '--estimators' in explicit else None
    requested_methods = [m.resolve_optimizer_name(o) for o in local.optimizers] if '--optimizers' in explicit else None
    results, indexed, sources, origins = {}, {}, {}, {}
    datasets, classifiers, algorithms = [], [], []
    rejected = []
    legacy_parameters = None
    kinds = ('results',) if local.experiment_mode == 'full' else ('results', 'progress')
    paths = [p for kind in kinds for p in sorted(cache.glob(f'{tag}_*_{kind}.pkl'))
             if '_snapshot_' not in p.name]
    if config is not None:
        preferred = [safe_path(manifest_path(local).parent / entry['path']) for entry in config['cache_signatures']
                     if '_snapshot_' not in entry['path'] and
                     (local.experiment_mode != 'full' or entry['path'].endswith('_results.pkl'))]
        if preferred and all(path.is_file() for path in preferred):
            paths = preferred
        else:
            print('[config] Some recorded cache paths are missing; discovering compatible existing caches', flush=True)
    for path in paths:
        path = safe_path(path)
        digest = sha256(path)
        if config is not None:
            entry = next((entry for entry in config['cache_signatures']
                          if manifest_path(local).parent / entry['path'] == path), None)
            if entry is not None and digest != entry['sha256']:
                raise ValueError(f'Manifest MISMATCH cache SHA256: {path}')
        try:
            payload = m.load_cache(str(path))
            if not isinstance(payload, dict) or not payload or not all(isinstance(r, dict) for r in payload.values()):
                raise ValueError('Invalid cache payload')
            estimators = {str(row.get('Estimator', '')).lower() for row in payload.values()}
            if len(estimators) != 1 or not all(estimators):
                raise ValueError('Ambiguous estimators in cache')
            classifier = estimators.pop()
            stem = path.stem[len(tag) + 1:].rsplit('_', 1)[0]
            dataset, separator, signature = stem.rpartition(f'_{classifier}_')
            if not separator or not dataset or not signature:
                raise ValueError('Invalid cache filename')
            if ((requested_datasets is not None and dataset not in requested_datasets)
                    or (requested_classifiers is not None and classifier not in requested_classifiers)):
                continue
            groups = {label: _report_algorithm(label, classifier) for label in payload}
            methods = list(dict.fromkeys(method for group, method in groups.values()))
            evidence = None
            if any('CacheIdentity' not in row for row in payload.values()):
                evidence = _legacy_report_identity(local, signature, methods)
        except (ValueError, TypeError, EOFError, pickle.UnpicklingError) as exc:
            rejected.append(f'{path.name}: {exc}')
            _audit_cache('MISMATCH', path.name, 'compatible completed cache', str(exc))
            continue
        accepted = False
        for label, row in payload.items():
            algorithm, method = groups[label]
            if requested_methods is not None and method not in requested_methods:
                continue
            try:
                _validate_report_row(row, local, dataset, classifier, method, evidence)
            except (ValueError, TypeError) as exc:
                rejected.append(f'{path.name}/{label}: {exc}')
                _audit_cache('MISMATCH', f'{path.name}/{label}', 'compatible completed row', str(exc))
                continue
            key = dataset, classifier, algorithm
            if key in indexed:
                if not _same_report_observations(indexed[key], row):
                    message = f'Conflicting compatible observations for {key}: {origins[key]} and {path.name}'
                    _audit_cache('MISMATCH', 'observations', 'identical duplicate', message)
                    raise ValueError(message)
                continue
            indexed[key] = row
            origins[key] = path.name
            accepted = True
            results.setdefault(dataset, {})[label] = row
            for value, collection in ((dataset, datasets), (classifier, classifiers), (algorithm, algorithms)):
                if value not in collection:
                    collection.append(value)
        if sha256(path) != digest:
            raise ValueError(f'Cache changed during read: {path}')
        if accepted:
            if evidence:
                contract = local.report_legacy_parameters
                if legacy_parameters is not None and legacy_parameters != contract:
                    raise ValueError('MISMATCH: mixed legacy parameter configurations in selected caches')
                legacy_parameters = contract
            sources[str(path.relative_to(root))] = digest
            _audit_cache('MATCH', path.name, f'{dataset}/{classifier}', f'{local.runs} completed runs',
                         evidence or 'CacheIdentity verified; filename hash ignored')
    if not indexed:
        for field, requested in (('datasets', requested_datasets), ('estimators', requested_classifiers),
                                 ('optimizers', requested_methods)):
            _audit_cache('MISMATCH', field, requested or 'discover compatible stored selection', [])
        for field in REPORT_SCIENCE_FIELDS:
            _audit_cache('MISMATCH', field, getattr(local, field), 'no verified compatible cache')
        if rejected:
            raise ValueError('Cache MISMATCH; no compatible completed observations: ' + '; '.join(rejected))
        raise FileNotFoundError(f'No completed caches in {cache}; no recomputation allowed')
    checks = [('datasets', requested_datasets or datasets, datasets),
              ('estimators', requested_classifiers or classifiers, classifiers),
              ('optimizers', requested_methods or list(dict.fromkeys(_report_algorithm(a, '')[1] for a in algorithms)),
               list(dict.fromkeys(_report_algorithm(a, '')[1] for a in algorithms)))]
    mismatches = []
    for field, requested, stored in checks:
        match = set(requested) == set(stored)
        _audit_cache('MATCH' if match else 'MISMATCH', field, requested, stored,
                     'selected from cache' if field in ('estimators', 'optimizers') and f'--{field}' not in explicit else '')
        if not match:
            mismatches.append(f'{field}: missing {sorted(set(requested) - set(stored))}')
    for field in REPORT_SCIENCE_FIELDS:
        _audit_cache('MATCH', field, getattr(local, field), getattr(local, field), 'verified metadata/legacy digest')
    if mismatches or len(indexed) != len(datasets) * len(classifiers) * len(algorithms):
        _audit_cache('MISMATCH', 'grid', 'complete dataset/estimator/optimizer grid', list(indexed))
        raise ValueError('Cache MISMATCH: incomplete compatible grid; ' + '; '.join(mismatches + rejected))
    metrics = [m for m in METRICS if all(m.run_key in row for row in indexed.values())]
    if not metrics:
        raise ValueError('No complete cached metrics')
    if legacy_parameters is not None:
        local.report_legacy_parameters = legacy_parameters
    elif hasattr(local, 'report_legacy_parameters'):
        del local.report_legacy_parameters
    return CompletedReport(local, results, indexed, datasets, classifiers, algorithms, metrics,
                           'stored-final-caches', sources)


def run_figure_report(args):
    """Write only figures; cache, numerical results and scientific execution are protected."""
    from reporting import figures
    root = safe_path(args.output_root)
    destination_root = safe_path(getattr(args, 'report_output_root', None) or root)
    with report_stage('Read stored final caches'), report_guard(()):
        reports = [load_cached_figure_report(framework().clone_args_for_mode(args, mode))
                   for mode in dict.fromkeys(args.experiment_modes)]
    protected = {p: sha256(p) for p in (root / 'Results' / f'EXP{args.exp_id:03d}').rglob('*')
                 if p.is_file()}
    destinations = []
    result_destinations = []
    for report in reports:
        language = figures.figure_language(report)
        mode = report.args.experiment_mode + ('_esp' if language == 'es' else '')
        destination = safe_path(destination_root / 'Figures' / report.exp_tag / mode)
        results_destination = safe_path(destination_root / 'Results' / report.exp_tag / mode)
        # Create ancestors before the guard; all actual exports stay in Figures.
        if report.args.experiment_mode == 'full':
            destination.parent.mkdir(parents=True, exist_ok=True)
            replica = 0
            while True:
                try:
                    destination.mkdir()
                    break
                except FileExistsError:
                    replica += 1
                    destination = safe_path(destination.parent / f'{mode}_rep{replica}')
        else:
            destination.mkdir(parents=True, exist_ok=True)
        print(f'[report] output folder = {destination.name}', flush=True)
        results_destination.mkdir(parents=True, exist_ok=True)
        with report_stage(f'Figures: {destination}'), report_guard((destination,)):
            figures.generate(report, destination)
            figures.generate_statistics_figures(report, destination)
        destinations.append(str(destination))
        result_destinations.append(str(results_destination))
    if any(sha256(path) != digest for path, digest in protected.items()):
        raise ValueError('Protected cache/results changed during figure reporting')
    from reporting.experiment_config import save_report_config
    for report in reports:
        save_report_config(report)
    print('[report] optimization runs=0', flush=True)
    return {'figures_destination': destinations, 'results_destination': result_destinations, 'optimization_calls': 0,
            'protected_files_unchanged': True}
